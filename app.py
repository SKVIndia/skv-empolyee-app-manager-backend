#!/usr/bin/env python3
import os
import csv
import io
import openpyxl
import pg8000
from flask import Flask, request, jsonify
from flask_cors import CORS
from dotenv import load_dotenv
from werkzeug.utils import secure_filename

# Load environment variables
load_dotenv()

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = "./uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

DATABASE_URL = os.getenv(
    "DATABASE_URL",
    "postgresql://neondb_owner:npg_yliZ19YbeQhV@ep-mute-cloud-a1ecwvbi-pooler.ap-southeast-1.aws.neon.tech/skv-employees?sslmode=require&channel_binding=require"
)

PORT = int(os.getenv("PORT", 5000))


@app.route("/", methods=["GET"])
def home():
    return jsonify({"message": "✅ SKV Employee Uploader is running."})


@app.route("/upload", methods=["POST"])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    file = request.files['file']
    if file.filename == "":
        return jsonify({"error": "Empty file name"}), 400

    filename = secure_filename(file.filename)
    file_path = os.path.join(UPLOAD_FOLDER, filename)
    file.save(file_path)

    try:
        create_neon_database(file_path)
        return jsonify({"message": "✅ Upload and database sync successful"}), 200
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


def create_neon_database(file_path: str):
    ext = file_path.split('.')[-1].lower()

    data = []
    headers = []

    if ext == 'csv':
        with open(file_path, newline='', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            headers = [h.strip().lower().replace(" ", "_") for h in reader.fieldnames]
            for row in reader:
                data.append({k.strip().lower().replace(" ", "_"): v.strip() for k, v in row.items()})
    elif ext in ['xlsx', 'xls']:
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        headers = [str(cell.value).strip().lower().replace(" ", "_") for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            data.append({headers[i]: (str(cell).strip() if cell is not None else None) for i, cell in enumerate(row)})
    else:
        raise ValueError("❌ Unsupported file format")

    if not data:
        raise ValueError("❌ File is empty or invalid")

    user_cols = ['employee_email', 'password']
    app_cols = [col for col in headers if col not in user_cols]

    conn = pg8000.connect(dsn=DATABASE_URL)
    cur = conn.cursor()

    cur.execute("DROP TABLE IF EXISTS permissions;")
    cur.execute("DROP TABLE IF EXISTS users CASCADE;")

    cur.execute(f"""
        CREATE TABLE users (
            employee_email TEXT PRIMARY KEY,
            password TEXT NOT NULL,
            {', '.join([f"{col} TEXT" for col in app_cols])}
        );
    """)

    cur.execute("""
        CREATE TABLE permissions (
            email TEXT REFERENCES users(employee_email) ON DELETE CASCADE,
            app_name TEXT NOT NULL,
            app_link TEXT
        );
    """)

    for row in data:
        values = [row.get(col) for col in headers]
        placeholders = ', '.join(['%s'] * len(headers))
        cur.execute(f"""
            INSERT INTO users ({', '.join(headers)})
            VALUES ({placeholders})
            ON CONFLICT (employee_email) DO UPDATE SET
            {', '.join([f"{col}=EXCLUDED.{col}" for col in headers if col != 'employee_email'])};
        """, values)

        email = row.get("employee_email")
        for app_col in app_cols:
            link = row.get(app_col)
            if link and link.lower() not in ["n/a", "none", "null", ""]:
                cur.execute("INSERT INTO permissions (email, app_name, app_link) VALUES (%s, %s, %s);", (email, app_col, link))

    conn.commit()
    cur.close()
    conn.close()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=PORT)
